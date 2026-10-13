"""Add Phase 1 to the supplied current cockpit, preserving every existing data block.
Usage: python scripts/build-daily-labour.py --cockpit HTML --crown XLSX --roster XLSX
"""
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime,date,time,timezone
import argparse,json,re,subprocess,hashlib,tempfile

ROOT=Path(__file__).resolve().parent.parent
p=argparse.ArgumentParser()
p.add_argument('--cockpit',required=True);p.add_argument('--crown',required=True);p.add_argument('--roster',required=True)
p.add_argument('--output',default=str(ROOT/'GC_NH_Cockpit_Daily_Labour_Phase1.html'))
a=p.parse_args()
def val(x):
    if isinstance(x,(datetime,date)):return x.isoformat()
    if isinstance(x,time):return x.isoformat()
    return x
def sheets(file,kind):
    wb=load_workbook(file,read_only=True,data_only=True)
    formulas=load_workbook(file,read_only=True,data_only=False) if kind=='crown' else None
    result=[]
    for s in wb:
        limit=18 if kind=='crown' else 9
        grid=[[val(v) for v in row[:limit]] for row in s.values]
        derived=False
        if formulas:
            for row in formulas[s.title].iter_rows(min_row=6,max_row=8):
                derived=derived or any(isinstance(c.value,str) and ('IF(P' in c.value or 'COUNTIF' in c.value) for c in row)
        result.append({'name':s.title,'grid':grid,'derivedDirection':derived})
    return result
inputs={'crown':sheets(a.crown,'crown'),'roster':sheets(a.roster,'roster')}
for kind in ['crown','roster']:
    f=Path(getattr(a,kind));inputs[kind+'Meta']={'id':kind+'-initial-13-10-2026','file':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'loadedAt':datetime.now(timezone.utc).isoformat(timespec='seconds')}
with tempfile.TemporaryDirectory() as tmp:
    input_path=Path(tmp)/'inputs.json';input_path.write_text(json.dumps(inputs))
    script="""const fs=require('fs'),E=require(process.argv[1]),x=JSON.parse(fs.readFileSync(process.argv[2]));
const roster=E.parseRoster(x.roster,x.rosterMeta,'ABC'),crown=E.parseCrown(x.crown,x.crownMeta);crown.meta.productionDate='2026-10-13';
const model={schema:1,rosters:{ABC:roster},rosterVersions:{[roster.meta.id]:roster},dayRosters:{'2026-10-13':{ABC:roster.meta.id}},crownFiles:[crown],overrides:{},employees:{},prep:{},history:{},audit:[],future:{efficiency:[],official:[]},ui:{date:'2026-10-13',unit:'ABC',search:'',exceptionScope:'unit'}};
const r=E.calculate(model,'2026-10-13');model.history['2026-10-13']={date:r.W.date,window:{from:r.W.from,to:r.W.to},areas:r.areas,coverage:r.coverage,exceptionCount:r.exceptions.length,status:'FLASH',outputUnits:null,official:null};
fs.writeFileSync(process.argv[3],JSON.stringify(model));fs.writeFileSync(process.argv[4],JSON.stringify({roster:roster.meta,crown:crown.meta,areas:r.areas,coverage:r.coverage,exceptionTypes:Object.fromEntries([...new Set(r.exceptions.map(e=>e.type))].map(t=>[t,r.exceptions.filter(e=>e.type===t).length])),rosterAudit:E.rosterAudit(null,roster),unknownRoles:[...new Set(roster.rows.filter(r=>E.classify(r).class==='REVIEW').map(r=>r.role))],identities:r.exceptions.filter(e=>e.type.includes('name')||e.type.includes('Duplicate Clock'))},null,2));"""
    seed=Path(tmp)/'seed.json';report=ROOT/'tests/Daily_Labour_Source_Inspection.json'
    subprocess.run(['node','-e',script,str(ROOT/'scripts/daily-labour/engine.js'),str(input_path),str(seed),str(report)],check=True)
    seed_text=seed.read_text()
source=Path(a.cockpit).read_text()
if 'id="daily-labour-module"' in source:raise SystemExit('Daily Labour is already installed. Use the existing module.')
# Exact script-block matching avoids matching inert-script literals inside JS.
def data_blocks(s):
    return {re.search(r'\bid=[\"\']([^\"\']+)',m[1])[1]:m[0] for m in re.finditer(r'<script\b([^>]*)>([\s\S]*?)</script\s*>',s,re.I) if re.search(r'\btype=[\"\']application/',m[1],re.I) and re.search(r'\bid=',m[1])}
before=data_blocks(source)
# Carry forward the user's separately requested daily complaint-date fix.
with tempfile.TemporaryDirectory() as tmp:
    f=Path(tmp)/'source.html';f.write_text(source);out=Path(tmp)/'patched.html'
    js="const fs=require('fs');eval(fs.readFileSync(process.argv[1],'utf8'));fs.writeFileSync(process.argv[3],applyDailyComplaintLag(fs.readFileSync(process.argv[2],'utf8')).html);"
    subprocess.run(['node','-e',js,str(ROOT/'scripts/quality/daily-complaint-lag-patch.js'),str(f),str(out)],check=True)
    source=out.read_text()
old="{ id: 'opp', label: 'Operations & People', tabs: ['ops', 'people'] }"
assert source.count(old)==1;source=source.replace(old,"{ id: 'opp', label: 'Operations & People', tabs: ['ops', 'people', 'daily_labour'] }")
old='    await QH.boot();';assert source.count(old)==1
source=source.replace(old,'    await DailyLabour.init();\n'+old)
# Additional modules can move PO outside the old 128 KB tail search. Preserve
# its protection using an inert HTML parse when the fast tail check misses it.
old="if(poTail.indexOf('<script id=\"po-module\">')<0)throw new Error"
new="if(poTail.indexOf('<script id=\"po-module\">')<0){var poDocument=new DOMParser().parseFromString(await file.text(),'text/html'),poScript=poDocument.querySelector('script#po-module');if(!poScript||!poScript.textContent.trim())throw new Error"
assert source.count(old)==1;source=source.replace(old,new)
old="This earlier file does not contain PO management.');}"
assert source.count(old)==1;source=source.replace(old,"This earlier file does not contain PO management.');}}")
css=(ROOT/'scripts/daily-labour/module.css').read_text()
insert='\n<style id="daily-labour-style">'+css+'</style>\n<script id="daily-labour-seed" type="application/json">'+seed_text.replace('<','\\u003c')+'</script>\n'
for name,id in [('engine','daily-labour-engine'),('module-phase1','daily-labour-module')]:
    js=(ROOT/f'scripts/daily-labour/{name}.js').read_text();assert '</script' not in js.lower()
    insert+='<script id="'+id+'">'+js+'</script>\n'
at=source.rindex('</body>');source=source[:at]+insert+source[at:]
after=data_blocks(source)
assert all(after.get(k)==v for k,v in before.items()),'Existing data blocks changed'
Path(a.output).write_text(source)
print(json.dumps({'output':a.output,'bytes':len(source.encode()),'originalDataBlocksPreserved':len(before),'seedBytes':len(seed_text),'sha256':hashlib.sha256(source.encode()).hexdigest()},indent=2))
