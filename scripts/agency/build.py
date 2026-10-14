"""Build the portable agency companion from user-supplied allocation workbooks."""
from pathlib import Path
import argparse,collections,datetime,hashlib,json,re
import openpyxl
ap=argparse.ArgumentParser();ap.add_argument('--jt',required=True);ap.add_argument('--quest',required=True);args=ap.parse_args()
root=Path(__file__).resolve().parents[2]; records=[]; sources=[]; invalid=[]
for agency, filename in [('J&T',args.jt),('Quest',args.quest)]:
 p=Path(filename);w=openpyxl.load_workbook(p,read_only=True,data_only=True);s=w['Data'];headers=None;date_cells=[]
 for row in s.iter_rows():
  vals=[c.value for c in row]
  if headers is None:
   if 'Date' in vals:
    headers=vals;ix={str(h):i for i,h in enumerate(vals) if h is not None}
   continue
  if not any(v is not None for v in vals):continue
  get=lambda k:vals[ix[k]] if k in ix else None
  d=get('Date');name=str(get('Full name') or '').strip();rn=row[0].row
  if not isinstance(d,datetime.datetime) or not name:
   invalid.append({'agency':agency,'sheet':'Data','row':rn,'reason':'Missing date or name'});continue
  role=str(get('L') if agency=='J&T' else get('Function/Line') or '').strip(); role='' if role=='None' else role
  unit=str(get('Unit') or '').strip();shift=str(get('Shift') or '').strip();raw_id=get('Employee Number')
  emp=str(raw_id or '').strip();emp=emp[:-2] if emp.endswith('.0') else emp
  ident=emp if re.fullmatch(r'\d+',emp) else 'name:'+re.sub(r'\s+',' ',name).casefold()
  kind=str(get('Direct / Indirect') or '').strip();san=bool(re.search(r'sanitis|sanitiz',role+' '+unit,re.I));mixed=san and role.casefold() not in ('sanitiser','sanitizer')
  cl='IDL' if san or kind.casefold()=='indirect' else 'DL' if kind.casefold()=='direct' else 'REVIEW'
  records.append({'agency':agency,'date':d.date().isoformat(),'name':name,'employee':emp,'identity':ident,'unit':unit,'location':role,'shift':shift,'class':cl,'mixedSanitiser':mixed,'sourceClass':kind,'row':rn,'dateFormat':row[ix['Date']].number_format,'dateAnomaly':d.year<2020 or d.year>2026})
 for sh in w:
  if sh.title in ('Data','Budget','List','JT FULL LIST','Report','Sheet1'):continue
  for row in sh.iter_rows(min_row=1,max_row=min(4,sh.max_row)):
   for c in row:
    if isinstance(c.value,datetime.datetime):date_cells.append({'sheet':sh.title,'cell':c.coordinate,'date':c.value.isoformat(),'numberFormat':c.number_format})
 sources.append({'agency':agency,'filename':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'printedDateCells':date_cells});w.close()
groups=collections.defaultdict(list)
for r in records:groups[(r['date'],r['identity'])].append(r)
people=[]
for (date,identity),rows in groups.items():
 locations={(r['unit'].casefold(),r['location'].casefold()) for r in rows};classes={r['class'] for r in rows};agencies=sorted({r['agency'] for r in rows});flags=[]
 if len(locations)>1:flags.append('Conflicting locations; 8 hours remain unallocated')
 if len(agencies)>1:flags.append('Appears in both agencies; counted once')
 if any(r['dateAnomaly'] for r in rows):flags.append('Source date requires review; preserved unchanged')
 if any(r['mixedSanitiser'] for r in rows):flags.append('Mixed Sanitiser location; provisional IDL pending Labour Report')
 if identity.startswith('name:'):flags.append('No valid employee number; matched on name')
 people.append({'date':date,'identity':identity,'name':rows[0]['name'],'agency':' / '.join(agencies),'unit':rows[0]['unit'] if len(locations)==1 else 'Review','location':rows[0]['location'] if len(locations)==1 else ' / '.join(sorted({r['unit']+' '+r['location'] for r in rows})),'class':next(iter(classes)) if len(classes)==1 and len(locations)==1 else 'REVIEW','hours':8,'shift':' / '.join(sorted({r['shift'] for r in rows})),'flags':flags,'sourceRows':rows})
people.sort(key=lambda r:(r['date'],r['name'].casefold(),r['identity']))
data={'schema':1,'assumption':'8 estimated hours per unique person per source date until the Labour Report is available. Not actual worked or paid hours.','classification':'Sanitiser estimates are IDL, including mixed Sanitiser location labels provisionally; conflicting locations remain REVIEW. Other classes follow an explicit source Direct / Indirect value; otherwise REVIEW.','sources':sources,'excludedRows':invalid,'people':people,'allocationRows':len(records)}
encoded=json.dumps(data,separators=(',',':'),ensure_ascii=False).replace('<','\\u003c')
template=(root/'scripts/agency/template.html').read_text();(root/'Agency_Allocations_Provisional.html').write_text(template.replace('__DATA__',encoded))
(root/'tests/Agency_Source_Verification.json').write_text(json.dumps({'sourceHashes':{s['filename']:s['sha256'] for s in sources},'allocationRows':len(records),'uniquePersonDays':len(people),'excludedRows':len(invalid),'sourceDatesPreserved':True,'hoursPerPersonDay':8,'actualHours':False},indent=2)+'\n')
print('Built',len(records),'allocations;',len(people),'unique person-days;',len(invalid),'excluded rows')
